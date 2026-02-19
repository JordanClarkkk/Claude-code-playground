"""Read, merge, and write Excel catalogs with openpyxl."""

import io
import copy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------

def read_catalog(excel_bytes: bytes) -> tuple[list[str], list[dict]]:
    """Read an Excel file and return (columns, rows_as_dicts)."""
    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [str(h).strip() if h is not None else f"Column_{i}"
               for i, h in enumerate(rows[0], 1)]
    data = []
    for row in rows[1:]:
        d = {}
        for h, v in zip(headers, row):
            if v is not None:
                d[h] = v
        if d:
            data.append(d)
    return headers, data


# ---------------------------------------------------------------------------
# Merging
# ---------------------------------------------------------------------------

def _normalise(val: object) -> str:
    """Lower-case, stripped string for fuzzy matching."""
    return str(val).strip().lower() if val is not None else ""


def _build_key(row: dict, key_columns: list[str]) -> str | None:
    """Build a composite match key from the row."""
    parts = [_normalise(row.get(c)) for c in key_columns]
    parts = [p for p in parts if p]
    return "|".join(parts) if parts else None


def _guess_key_columns(columns: list[str]) -> list[str]:
    """Heuristically pick columns that look like product identifiers."""
    priority_keywords = ["sku", "part", "upc", "ean", "barcode", "code", "id", "item"]
    name_keywords = ["product", "name", "title", "description"]

    keys: list[str] = []
    for kw in priority_keywords:
        for col in columns:
            if kw in col.lower() and col not in keys:
                keys.append(col)
    if not keys:
        for kw in name_keywords:
            for col in columns:
                if kw in col.lower() and col not in keys:
                    keys.append(col)
    return keys


def merge_products(
    existing_columns: list[str],
    existing_rows: list[dict],
    new_products: list[dict],
) -> tuple[list[str], list[dict], dict]:
    """Merge new_products into the existing catalog.

    Returns (final_columns, merged_rows, stats).
    stats = {"matched": int, "added": int, "new_columns": list[str]}
    """
    # Discover any new columns
    all_new_keys: set[str] = set()
    for p in new_products:
        all_new_keys.update(p.keys())
    new_columns = [k for k in all_new_keys if k not in existing_columns]
    final_columns = existing_columns + sorted(new_columns)

    key_columns = _guess_key_columns(final_columns)

    # Index existing rows
    index: dict[str, int] = {}
    if key_columns:
        for i, row in enumerate(existing_rows):
            k = _build_key(row, key_columns)
            if k:
                index[k] = i

    merged = [dict(r) for r in existing_rows]  # shallow copy each row
    matched = 0
    added = 0

    for prod in new_products:
        k = _build_key(prod, key_columns) if key_columns else None
        if k and k in index:
            # Update existing row — only fill in blanks or add new columns
            row_idx = index[k]
            for col, val in prod.items():
                if col not in merged[row_idx] or merged[row_idx][col] in (None, ""):
                    merged[row_idx][col] = val
            matched += 1
        else:
            merged.append(prod)
            added += 1

    stats = {
        "matched": matched,
        "added": added,
        "new_columns": new_columns,
    }
    return final_columns, merged, stats


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_NEW_COL_FILL = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def write_catalog(
    columns: list[str],
    rows: list[dict],
    new_columns: list[str] | None = None,
) -> bytes:
    """Write a styled Excel workbook and return it as bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Product Catalog"

    new_col_set = set(new_columns or [])

    # Headers
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _NEW_COL_FILL if col_name in new_col_set else _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    # Data
    alt_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, col_name in enumerate(columns, 1):
            val = row_data.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center")
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Auto-width
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(str(col_name))
        for row_idx in range(2, len(rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 4, 50)

    # Freeze header row
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
